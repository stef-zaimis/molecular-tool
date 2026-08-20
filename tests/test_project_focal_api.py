"""
The focal-set service API, duplicate-header refusal, and run gating.

These are the rules the renderer is about to depend on, tested at the service
boundary rather than through React, because a rule that only exists as a
disabled button is not a rule about the data.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from molecular_diagnosis.project.indexing import scan_fasta
from molecular_diagnosis.project.locations import PresenceState
from molecular_diagnosis.project.service import (
    ProjectError,
    ProjectService,
    parse_focal_text,
)

RECORDS = [
    ("focal_1|AU|Target", "AACC"),
    ("focal_2|GB|Target", "AACC"),
    ("other_1|FR|Contrast", "AGCT"),
    ("other_2|DE|Contrast", "CGTT"),
]


def write_fasta(path, records):
    path.write_text(
        "".join(f">{header}\n{sequence}\n" for header, sequence in records), encoding="utf-8"
    )
    return path


@pytest.fixture()
def project(tmp_path):
    service = ProjectService(tmp_path / "proj")
    service.open_project("Test")
    yield service
    service.close()


@pytest.fixture()
def fasta(tmp_path):
    return write_fasta(tmp_path / "a.fasta", RECORDS)


@pytest.fixture()
def linked(project, fasta):
    row, _status = project.link_fasta(str(fasta))
    return row


@pytest.fixture()
def focal_set(project):
    return project.create_focal_set("Targets")


def outputs_written(project) -> list[str]:
    return sorted(path.name for path in project.outputs_dir.iterdir())


# ---------------------------------------------------------------------------
# 1. Duplicate headers inside one FASTA
# ---------------------------------------------------------------------------


DUPLICATE_RECORDS = [
    ("dupe|Target", "AACC"),
    ("dupe|Target", "AAGG"),
    ("other_1|FR|Contrast", "AGCT"),
    ("other_2|DE|Contrast", "CGTT"),
]


def test_a_file_with_repeated_headers_still_indexes(project, tmp_path):
    """The file stays linkable and searchable; only ANALYSIS is refused."""
    path = write_fasta(tmp_path / "dupes.fasta", DUPLICATE_RECORDS)
    row, status = project.link_fasta(str(path))

    assert status.available
    assert project.repository.record_count(row.id) == 4
    assert row.duplicate_header_count == 1
    assert project.search_headers("dupe")["hits"]


def test_single_file_run_is_refused_when_the_file_has_duplicate_headers(project, tmp_path):
    path = write_fasta(tmp_path / "dupes.fasta", DUPLICATE_RECORDS)
    row, _status = project.link_fasta(str(path))
    focal = project.create_focal_set("S")
    project.replace_focal_entries(focal["id"], headers=["dupe|Target"])

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal["id"], fasta_file_ids=[row.id], single_file=True, options={}
        )

    assert caught.value.code == "DUPLICATE_HEADER_IN_FILE"
    assert "dupes.fasta" in (caught.value.detail or "")
    assert "1 duplicate" in (caught.value.detail or "")


def test_all_files_run_is_refused_when_one_member_has_duplicate_headers(
    project, fasta, linked, tmp_path
):
    dupes = write_fasta(
        tmp_path / "dupes.fasta", [("d1|X", "AAAA"), ("d1|X", "TTTT"), ("d2|X", "GGGG")]
    )
    dupes_row, _status = project.link_fasta(str(dupes))
    focal = project.create_focal_set("S")
    project.add_focal_entries_from_query(focal["id"], "Target")

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal["id"],
            fasta_file_ids=[linked.id, dupes_row.id],
            single_file=False,
            options={},
        )

    assert caught.value.code == "DUPLICATE_HEADER_IN_FILE"
    assert "dupes.fasta" in (caught.value.detail or "")


def test_no_sequence_is_silently_dropped_or_analysed(project, tmp_path):
    """
    The refusal exists because the alternative is invisible data loss.

    `scan_fasta` keeps both records, but the header-keyed `sequences` dict
    keeps one — so an analysis would quietly study "AAGG" and never mention
    that "AACC" existed. Nothing is renamed, nothing is dropped, and no output
    is produced.
    """
    path = write_fasta(tmp_path / "dupes.fasta", DUPLICATE_RECORDS)
    scan = scan_fasta(path)
    assert len(scan.records) == 4
    assert len(scan.sequences) == 3
    assert scan.duplicate_header_count == 1

    row, _status = project.link_fasta(str(path))
    focal = project.create_focal_set("S")
    project.replace_focal_entries(focal["id"], headers=["dupe|Target"])

    with pytest.raises(ProjectError):
        project.run_molecular_diagnosis(
            focal_set_id=focal["id"], fasta_file_ids=[row.id], single_file=True, options={}
        )

    assert outputs_written(project) == []
    # Both records survive in the index: refusing to analyse is not deleting.
    assert project.repository.record_count(row.id) == 4


# ---------------------------------------------------------------------------
# 2. Focal-set service API and locking
# ---------------------------------------------------------------------------


def test_the_focal_set_lifecycle(project, linked):
    created = project.create_focal_set("Targets")
    assert created["title"] == "Targets"
    assert created["locked"] is False
    assert created["entries"] == []

    renamed = project.rename_focal_set(created["id"], "Renamed")
    assert renamed["title"] == "Renamed"

    assert [s["id"] for s in project.list_focal_sets()] == [created["id"]]

    project.delete_focal_set(created["id"])
    assert project.list_focal_sets() == []


def test_an_unknown_focal_set_id_produces_a_project_error_not_a_sqlite_one(project):
    for call in (
        lambda: project.get_focal_set("nope"),
        lambda: project.rename_focal_set("nope", "x"),
        lambda: project.delete_focal_set("nope"),
        lambda: project.set_focal_set_locked("nope", True),
        lambda: project.replace_focal_entries("nope", headers=["h"]),
        lambda: project.add_focal_entries_from_query("nope", "Target"),
        lambda: project.remove_focal_entries_by_query("nope", "Target"),
        lambda: project.focal_presence("nope", selected_file_id=None),
        lambda: project.focal_selector("nope"),
    ):
        with pytest.raises(ProjectError) as caught:
            call()
        assert caught.value.code == "UNKNOWN_FOCAL_SET"


def test_a_locked_focal_set_refuses_every_mutation(project, linked, focal_set):
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    locked = project.set_focal_set_locked(focal_set["id"], True)
    assert locked["locked"] is True

    for call in (
        lambda: project.rename_focal_set(focal_set["id"], "New name"),
        lambda: project.delete_focal_set(focal_set["id"]),
        lambda: project.add_focal_entries_from_query(focal_set["id"], "Contrast"),
        lambda: project.remove_focal_entries_by_query(focal_set["id"], "focal_1"),
        lambda: project.replace_focal_entries(focal_set["id"], text="anything"),
    ):
        with pytest.raises(ProjectError) as caught:
            call()
        assert caught.value.code == "FOCAL_SET_LOCKED"

    # Nothing got through.
    stored = [entry.header for entry in project.repository.list_focal_entries(focal_set["id"])]
    assert stored == ["focal_1|AU|Target", "focal_2|GB|Target"]


def test_a_locked_focal_set_can_still_be_read_checked_and_analysed(project, linked, focal_set):
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    project.set_focal_set_locked(focal_set["id"], True)

    assert list(project.focal_selector(focal_set["id"])) == [
        "focal_1|AU|Target",
        "focal_2|GB|Target",
    ]
    presence = project.focal_presence(focal_set["id"], selected_file_id=linked.id)
    assert all(entry.state is PresenceState.PRESENT_CURRENT for entry in presence)

    result = project.run_molecular_diagnosis(
        focal_set_id=focal_set["id"],
        fasta_file_ids=[linked.id],
        single_file=True,
        options={},
    )
    assert result["focalHeaders"] == ["focal_1|AU|Target", "focal_2|GB|Target"]


def test_unlocking_is_the_one_mutation_a_locked_set_allows(project, linked, focal_set):
    project.set_focal_set_locked(focal_set["id"], True)
    assert project.set_focal_set_locked(focal_set["id"], False)["locked"] is False
    project.rename_focal_set(focal_set["id"], "Editable again")


# ---------------------------------------------------------------------------
# 3. Atomic explicit replacement (the large textbox)
# ---------------------------------------------------------------------------


def test_textbox_parsing_semantics():
    assert parse_focal_text("header one ; header two;header three") == [
        "header one",
        "header two",
        "header three",
    ]
    # Empty entries are ignored, not stored as empty headers.
    assert parse_focal_text(" ; ;;a;; ") == ["a"]
    assert parse_focal_text("") == []
    # Set semantics: an exact duplicate collapses onto its first occurrence.
    assert parse_focal_text("a;b;a") == ["a", "b"]
    # Whitespace INSIDE a header is part of the header; FASTA headers have spaces.
    assert parse_focal_text("gi|123 Homo sapiens") == ["gi|123 Homo sapiens"]


def test_replacement_stores_exactly_the_textbox_contents_in_order(project, linked, focal_set):
    result = project.replace_focal_entries(
        focal_set["id"], text="other_2|DE|Contrast ; focal_1|AU|Target"
    )

    assert result["headers"] == ["other_2|DE|Contrast", "focal_1|AU|Target"]
    stored = [entry.header for entry in project.repository.list_focal_entries(focal_set["id"])]
    assert stored == ["other_2|DE|Contrast", "focal_1|AU|Target"]


def test_replacement_keeps_ids_and_locations_of_unchanged_entries(project, linked, focal_set):
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    before = {e.header: e.id for e in project.repository.list_focal_entries(focal_set["id"])}
    kept_id = before["focal_1|AU|Target"]
    locations_before = project.repository.locations_for_entry(kept_id)
    assert locations_before

    result = project.replace_focal_entries(
        focal_set["id"], text="focal_1|AU|Target; other_1|FR|Contrast"
    )

    assert result["kept"] == ["focal_1|AU|Target"]
    assert result["added"] == ["other_1|FR|Contrast"]
    assert result["removed"] == ["focal_2|GB|Target"]

    after = {e.header: e.id for e in project.repository.list_focal_entries(focal_set["id"])}
    # The surviving entry is the SAME row: a debounced keystroke must not
    # destroy stable ids and their verified location cache.
    assert after["focal_1|AU|Target"] == kept_id
    assert project.repository.locations_for_entry(kept_id) == locations_before
    # The removed entry took its locations with it.
    assert project.repository.locations_for_entry(before["focal_2|GB|Target"]) == []
    # The new entry was seeded from the already-trusted index.
    assert project.repository.locations_for_entry(after["other_1|FR|Contrast"])


def test_repeated_identical_replacements_change_nothing(project, linked, focal_set):
    text = "focal_1|AU|Target; focal_2|GB|Target"
    project.replace_focal_entries(focal_set["id"], text=text)
    ids = {e.header: e.id for e in project.repository.list_focal_entries(focal_set["id"])}

    second = project.replace_focal_entries(focal_set["id"], text=text)

    assert second["added"] == [] and second["removed"] == []
    assert {
        e.header: e.id for e in project.repository.list_focal_entries(focal_set["id"])
    } == ids


def test_a_header_that_exists_in_no_fasta_is_kept_so_it_can_show_red(
    project, linked, focal_set
):
    project.replace_focal_entries(
        focal_set["id"], text="focal_1|AU|Target; typed_by_hand_and_wrong"
    )

    stored = [entry.header for entry in project.repository.list_focal_entries(focal_set["id"])]
    assert stored == ["focal_1|AU|Target", "typed_by_hand_and_wrong"]

    presence = {
        entry.header: entry.state
        for entry in project.focal_presence(focal_set["id"], selected_file_id=linked.id)
    }
    assert presence["focal_1|AU|Target"] is PresenceState.PRESENT_CURRENT
    assert presence["typed_by_hand_and_wrong"] is PresenceState.MISSING


def test_replacement_collapses_duplicates_preserving_the_first_occurrence(
    project, linked, focal_set
):
    result = project.replace_focal_entries(
        focal_set["id"], text="focal_1|AU|Target; focal_2|GB|Target; focal_1|AU|Target"
    )

    assert result["headers"] == ["focal_1|AU|Target", "focal_2|GB|Target"]
    assert [
        entry.header for entry in project.repository.list_focal_entries(focal_set["id"])
    ] == ["focal_1|AU|Target", "focal_2|GB|Target"]


def test_replacement_with_an_empty_textbox_clears_the_set(project, linked, focal_set):
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    result = project.replace_focal_entries(focal_set["id"], text="   ")

    # Unlike a `-` query, an emptied textbox is an explicit statement about
    # membership, so it is honoured.
    assert result["removed"] == ["focal_1|AU|Target", "focal_2|GB|Target"]
    assert project.repository.list_focal_entries(focal_set["id"]) == []


def test_replacement_requires_exactly_one_of_text_or_headers(project, focal_set):
    with pytest.raises(ProjectError) as caught:
        project.replace_focal_entries(focal_set["id"])
    assert caught.value.code == "INVALID_PARAMETER"

    with pytest.raises(ProjectError):
        project.replace_focal_entries(focal_set["id"], text="a", headers=["a"])


# ---------------------------------------------------------------------------
# 4. `+` and `-` semantics
# ---------------------------------------------------------------------------


def test_an_empty_add_query_is_refused_and_changes_nothing(project, linked, focal_set):
    for query in ("", "   ", "\t"):
        with pytest.raises(ProjectError) as caught:
            project.add_focal_entries_from_query(focal_set["id"], query)
        assert caught.value.code == "FOCAL_QUERY_EMPTY"
    assert project.repository.list_focal_entries(focal_set["id"]) == []


def test_an_empty_remove_query_does_not_erase_the_whole_set(project, linked, focal_set):
    """
    The regression this exists for.

    `"" in header` is true for every header, so an unguarded empty `-` query
    silently deleted the entire focal set.
    """
    project.add_focal_entries_from_query(focal_set["id"], "Target")

    for query in ("", "   "):
        with pytest.raises(ProjectError) as caught:
            project.remove_focal_entries_by_query(focal_set["id"], query)
        assert caught.value.code == "FOCAL_QUERY_EMPTY"

    assert [
        entry.header for entry in project.repository.list_focal_entries(focal_set["id"])
    ] == ["focal_1|AU|Target", "focal_2|GB|Target"]


def test_add_refuses_rather_than_persisting_a_partial_expansion(
    project, fasta, linked, focal_set, tmp_path
):
    """
    `+` promises every match in the requested scope.

    A file that cannot be read cannot be searched, so the honest answer is a
    refusal — not a smaller set that looks complete and is not reproducible.
    """
    other = write_fasta(tmp_path / "b.fasta", [("b1|Target", "AACC"), ("b2|X", "TTTT")])
    other_row, _status = project.link_fasta(str(other))
    other.unlink()

    with pytest.raises(ProjectError) as caught:
        project.add_focal_entries_from_query(focal_set["id"], "Target")

    assert caught.value.code == "SEARCH_SCOPE_UNAVAILABLE"
    assert "b.fasta" in (caught.value.detail or "")
    assert project.repository.list_focal_entries(focal_set["id"]) == []

    # Narrowing the scope to the file that IS readable is a different, complete
    # request, and is allowed.
    result = project.add_focal_entries_from_query(
        focal_set["id"], "Target", fasta_file_ids=[linked.id]
    )
    assert result["added"] == ["focal_1|AU|Target", "focal_2|GB|Target"]
    assert other_row.id


def test_search_reports_unusable_files_instead_of_refusing(
    project, fasta, linked, tmp_path
):
    other = write_fasta(tmp_path / "b.fasta", [("b1|Target", "AACC")])
    project.link_fasta(str(other))
    other.unlink()

    result = project.search_headers("Target")

    assert [hit["header"] for hit in result["hits"]] == [
        "focal_1|AU|Target",
        "focal_2|GB|Target",
    ]
    assert [entry["displayName"] for entry in result["unavailable"]] == ["b.fasta"]


def test_none_means_the_whole_project_and_empty_means_an_empty_scope(
    project, linked, focal_set
):
    whole = project.add_focal_entries_from_query(focal_set["id"], "Target", fasta_file_ids=None)
    assert whole["added"] == ["focal_1|AU|Target", "focal_2|GB|Target"]

    empty_set = project.create_focal_set("Empty scope")
    nothing = project.add_focal_entries_from_query(
        empty_set["id"], "Target", fasta_file_ids=[]
    )
    assert nothing["matched"] == [] and nothing["added"] == []
    assert project.repository.list_focal_entries(empty_set["id"]) == []


def test_remove_does_not_re_search_the_fasta(project, fasta, linked, focal_set):
    project.replace_focal_entries(focal_set["id"], headers=["focal_1|AU|Target"])
    # The file also contains focal_2, but it is not a member, so a `-` query
    # matching it removes nothing: `-` operates on the set, not on the file.
    result = project.remove_focal_entries_by_query(focal_set["id"], "focal_2")
    assert result["removed"] == []
    assert len(project.repository.list_focal_entries(focal_set["id"])) == 1


# ---------------------------------------------------------------------------
# 5. Focal completeness for every analysis scope
# ---------------------------------------------------------------------------


def test_all_files_run_is_refused_when_an_entry_is_in_none_of_the_files(
    project, fasta, linked, focal_set, tmp_path
):
    other = write_fasta(tmp_path / "b.fasta", [("b1|X", "AAAA"), ("b2|X", "TTTT")])
    other_row, _status = project.link_fasta(str(other))
    project.replace_focal_entries(
        focal_set["id"], headers=["focal_1|AU|Target", "never_anywhere|ZZ"]
    )

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal_set["id"],
            fasta_file_ids=[linked.id, other_row.id],
            single_file=False,
            options={},
        )

    assert caught.value.code == "FOCAL_ENTRIES_NOT_IN_SCOPE"
    assert "never_anywhere|ZZ" in (caught.value.detail or "")
    # The focal group was not quietly shrunk to the entry that did resolve.
    assert outputs_written(project) == []


def test_an_entry_existing_only_outside_the_scope_does_not_count(
    project, fasta, linked, focal_set, tmp_path
):
    """
    Presence must be judged against the files the run will actually read.

    `elsewhere|QQ|Target` is in the project, so a project-wide presence check
    calls it present — but it is not in the two files being analysed, and
    giving it credit for that would silently drop it from the focal group.
    """
    inside = write_fasta(tmp_path / "b.fasta", [("b1|X", "AAAA"), ("b2|X", "TTTT")])
    inside_row, _status = project.link_fasta(str(inside))
    outside = write_fasta(tmp_path / "c.fasta", [("elsewhere|QQ|Target", "AACC")])
    outside_row, _status = project.link_fasta(str(outside))

    project.add_focal_entries_from_query(focal_set["id"], "Target")
    assert len(project.repository.list_focal_entries(focal_set["id"])) == 3

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal_set["id"],
            fasta_file_ids=[linked.id, inside_row.id],
            single_file=False,
            options={},
        )

    assert caught.value.code == "FOCAL_ENTRIES_NOT_IN_SCOPE"
    assert "elsewhere|QQ|Target" in (caught.value.detail or "")

    # Widening the scope to include the file that holds it makes the same set
    # legitimate again.
    result = project.run_molecular_diagnosis(
        focal_set_id=focal_set["id"],
        fasta_file_ids=[linked.id, inside_row.id, outside_row.id],
        single_file=False,
        options={},
    )
    assert len(result["focalHeaders"]) == 3


def test_a_run_is_refused_when_a_required_source_is_unavailable(
    project, fasta, linked, focal_set, tmp_path
):
    missing = write_fasta(tmp_path / "b.fasta", [("b1|Target", "AACC"), ("b2|X", "TTTT")])
    missing_row, _status = project.link_fasta(str(missing))
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    missing.unlink()

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal_set["id"],
            fasta_file_ids=[linked.id, missing_row.id],
            single_file=False,
            options={},
        )

    # The entry that lived only in the vanished file cannot be established
    # either way, which is not the same as knowing it is absent.
    assert caught.value.code == "FOCAL_PRESENCE_UNKNOWN"
    assert "b1|Target" in (caught.value.detail or "")
    assert outputs_written(project) == []


def test_an_unavailable_scope_file_still_blocks_a_run_whose_focal_set_resolves(
    project, fasta, linked, focal_set, tmp_path
):
    missing = write_fasta(tmp_path / "b.fasta", [("b1|X", "AACC"), ("b2|X", "TTTT")])
    missing_row, _status = project.link_fasta(str(missing))
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    missing.unlink()

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal_set["id"],
            fasta_file_ids=[linked.id, missing_row.id],
            single_file=False,
            options={},
        )

    assert caught.value.code == "SOURCE_UNAVAILABLE"
    assert outputs_written(project) == []


def test_an_all_files_run_succeeds_when_every_entry_is_present(
    project, fasta, linked, focal_set, tmp_path
):
    other = write_fasta(tmp_path / "b.fasta", [("b1|ES|Target", "AACG"), ("b2|X", "TTTT")])
    other_row, _status = project.link_fasta(str(other))
    project.add_focal_entries_from_query(focal_set["id"], "Target")

    result = project.run_molecular_diagnosis(
        focal_set_id=focal_set["id"],
        fasta_file_ids=[linked.id, other_row.id],
        single_file=False,
        options={},
    )

    assert result["sequenceCount"] == 6
    assert result["focalHeaders"] == [
        "focal_1|AU|Target",
        "focal_2|GB|Target",
        "b1|ES|Target",
    ]
    assert outputs_written(project)


def test_presence_is_reported_against_the_requested_scope(
    project, fasta, linked, focal_set, tmp_path
):
    outside = write_fasta(tmp_path / "c.fasta", [("elsewhere|QQ|Target", "AACC")])
    outside_row, _status = project.link_fasta(str(outside))
    project.add_focal_entries_from_query(focal_set["id"], "Target")

    scoped = {
        entry.header: entry.state
        for entry in project.focal_presence(
            focal_set["id"], selected_file_id=None, fasta_file_ids=[linked.id]
        )
    }
    assert scoped["elsewhere|QQ|Target"] is PresenceState.MISSING

    project_wide = {
        entry.header: entry.state
        for entry in project.focal_presence(focal_set["id"], selected_file_id=None)
    }
    assert project_wide["elsewhere|QQ|Target"] is PresenceState.PRESENT_CURRENT
    assert outside_row.id


# ---------------------------------------------------------------------------
# 6. Exact membership, all the way through the pipeline and its outputs
# ---------------------------------------------------------------------------


def test_exact_membership_survives_into_the_report_and_the_workbook(project, tmp_path):
    """
    `ABC123` must not drag in `ABC123_extra`.

    Partitioning is tested elsewhere; this runs the real pipeline so the
    distinction is checked where it could actually be lost — in the report and
    the Excel writer, both of which re-derive focal membership from the
    selector rather than being handed the header list.
    """
    from openpyxl import load_workbook

    path = write_fasta(
        tmp_path / "prefix.fasta",
        [
            ("ABC123", "AACC"),
            ("ABC123_extra", "AACC"),
            ("other|X", "AGCT"),
            ("other2|X", "CGTT"),
        ],
    )
    row, _status = project.link_fasta(str(path))
    focal = project.create_focal_set("Exact")
    project.replace_focal_entries(focal["id"], headers=["ABC123"])

    result = project.run_molecular_diagnosis(
        focal_set_id=focal["id"], fasta_file_ids=[row.id], single_file=True, options={}
    )

    assert result["focalHeaders"] == ["ABC123"]

    text = Path(result["outputs"]["reportTxt"]).read_text(encoding="utf-8")
    assert "Focal sequences: 1\n" in text
    assert "Non-focal sequences: 3\n" in text

    # The workbook lists every NON-focal sequence plus the reference. If
    # membership had degraded to substring matching, ABC123_extra would have
    # been treated as focal and excluded from the sheet.
    sheet = load_workbook(result["outputs"]["workbookXlsx"])["Full"]
    ids = [cell.value for (cell,) in sheet.iter_rows(min_row=2, max_col=1)]
    assert "ABC123" in ids
    assert "ABC123_extra" in ids


def test_the_legacy_substring_path_is_unchanged(tmp_path):
    """The old non-project API still means substring containment."""
    from molecular_diagnosis.focal import partition_headers

    headers = ["ABC123", "ABC123_extra", "other"]
    focal, non_focal = partition_headers(headers, ["ABC123"])
    assert focal == ["ABC123", "ABC123_extra"]
    assert non_focal == ["other"]


def test_search_results_follow_the_projects_own_file_order(project, fasta, linked, tmp_path):
    """
    Regression: results were ordered by `fasta_file_id`, a random UUID.

    `+` stores the resolved headers in the order the search returned them, so
    ordering by a UUID meant the same query built a differently-ordered focal
    set on each run — and made a test asserting that order flaky rather than
    wrong.
    """
    second = write_fasta(tmp_path / "b.fasta", [("b1|Target", "AACC"), ("b2|X", "TTTT")])
    second_row, _status = project.link_fasta(str(second))

    project_order = [row.id for row in project.repository.list_fasta_files()]
    assert project_order == [linked.id, second_row.id]

    hits = project.search_headers("Target")["hits"]
    assert [hit["fastaFileId"] for hit in hits] == [linked.id, linked.id, second_row.id]
    assert [hit["header"] for hit in hits] == [
        "focal_1|AU|Target",
        "focal_2|GB|Target",
        "b1|Target",
    ]


def test_a_single_file_run_on_a_missing_file_says_so_rather_than_absent(
    project, fasta, linked, focal_set
):
    """
    Unreadable is not the same as absent.

    Reporting "that entry is not in the file" about a file nobody could read
    asserts something never observed, and sends the user looking for the wrong
    problem.
    """
    project.add_focal_entries_from_query(focal_set["id"], "Target")
    fasta.unlink()

    with pytest.raises(ProjectError) as caught:
        project.run_molecular_diagnosis(
            focal_set_id=focal_set["id"],
            fasta_file_ids=[linked.id],
            single_file=True,
            options={},
        )

    assert caught.value.code == "FOCAL_PRESENCE_UNKNOWN"
